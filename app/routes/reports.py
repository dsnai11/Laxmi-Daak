from flask import Blueprint, render_template, request, send_file, current_app, jsonify
from flask_login import login_required, current_user
from app import db
from app.models.parcel import Parcel, STATUS_LABELS
from app.models.branch import Branch
from app.models.scan_log import ScanLog
from app.models.user import User
from datetime import datetime, timedelta
from sqlalchemy import func
import openpyxl
from io import BytesIO

reports_bp = Blueprint('reports', __name__, url_prefix='/reports')

@reports_bp.before_request
@login_required
def check_reports_access():
    from flask import abort
    if current_user.role != 'superadmin' and current_user.role != 'manager' and not current_user.has_permission('can_reports'):
        abort(403)


@reports_bp.route('')
@reports_bp.route('/')
@login_required
def index():
    date_from = request.args.get('date_from', (datetime.utcnow() - timedelta(days=30)).strftime('%Y-%m-%d'))
    date_to = request.args.get('date_to', datetime.utcnow().strftime('%Y-%m-%d'))
    
    start_date = datetime.strptime(date_from, '%Y-%m-%d')
    end_date = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1) - timedelta(microseconds=1)
    
    query = Parcel.query.filter(Parcel.dispatched_at.between(start_date, end_date))
    
    stats = {
        'total': query.count(),
        'pending': query.filter(Parcel.status == 'pending').count(),
        'delivered': query.filter(Parcel.status == 'delivered').count(),
        'in_transit': query.filter(Parcel.status == 'in_transit').count()
    }
    
    # 1. Daily volume trend (Line chart)
    daily_stats = db.session.query(
        func.date(Parcel.dispatched_at).label('date'),
        func.count(Parcel.id).label('count')
    ).filter(
        Parcel.dispatched_at.between(start_date, end_date)
    ).group_by(
        func.date(Parcel.dispatched_at)
    ).order_by('date').all()
    
    trend_labels = [str(x[0]) for x in daily_stats]
    trend_data = [x[1] for x in daily_stats]

    # 2. Status distribution (Doughnut chart)
    status_stats = db.session.query(
        Parcel.status,
        func.count(Parcel.id)
    ).filter(
        Parcel.dispatched_at.between(start_date, end_date)
    ).group_by(Parcel.status).all()
    
    status_labels = [x[0].replace('_', ' ').title() for x in status_stats]
    status_data = [x[1] for x in status_stats]

    # 3. Branch volume (Bar chart)
    b_stats = db.session.query(
        Branch.name,
        func.count(Parcel.id)
    ).join(
        Parcel, Parcel.sender_branch_id == Branch.id
    ).filter(
        Parcel.dispatched_at.between(start_date, end_date)
    ).group_by(Branch.id).order_by(func.count(Parcel.id).desc()).limit(5).all()
    
    branch_labels = [x[0] for x in b_stats]
    branch_data = [x[1] for x in b_stats]
    
    branches = Branch.query.order_by(Branch.name).all()
    
    return render_template('reports/index.html', stats=stats, branches=branches, 
                           date_from=date_from, date_to=date_to,
                           trend_labels=trend_labels, trend_data=trend_data,
                           status_labels=status_labels, status_data=status_data,
                           branch_labels=branch_labels, branch_data=branch_data)

@reports_bp.route('/pending')
@login_required
def pending():
    branch_id = request.args.get('branch_id', type=int)
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    
    query = Parcel.query.filter(Parcel.status.in_(['pending', 'in_transit']))
    if branch_id:
        query = query.filter(Parcel.current_branch_id == branch_id)
    if date_from:
        query = query.filter(Parcel.dispatched_at >= datetime.strptime(date_from, '%Y-%m-%d'))
    if date_to:
        query = query.filter(Parcel.dispatched_at <= datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1) - timedelta(microseconds=1))
        
    parcels = query.all()
    
    return jsonify([{
        'tracking_number': p.tracking_number,
        'status': p.status,
        'current_branch': p.current_branch.name if p.current_branch else 'N/A',
        'created_at': p.dispatched_at.isoformat() if p.dispatched_at else None
    } for p in parcels])

@reports_bp.route('/export/excel')
@login_required
def export_excel():
    report_type = request.args.get('report_type', 'all')
    branch_id = request.args.get('branch_id', 'all')
    direction = request.args.get('direction', 'both')  # sent, received, both
    transfer_mode = request.args.get('transfer_mode', 'all') # by_hand, by_bus, by_courier, all
    courier_partner = request.args.get('courier_partner', 'all') # all or company name
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    
    query = Parcel.query
    
    # 1. Branch & Direction Filtering
    if not current_user.is_superadmin():
        # Regular user / manager: scope to their branch or personal account
        user_branch_id = current_user.branch_id
        if user_branch_id:
            if direction == 'sent':
                query = query.filter((Parcel.sender_branch_id == user_branch_id) | (Parcel.created_by == current_user.id))
            elif direction == 'received':
                query = query.filter((Parcel.receiver_branch_id == user_branch_id) | (Parcel.receiver_email == current_user.email))
            else: # both
                query = query.filter(
                    (Parcel.sender_branch_id == user_branch_id) | 
                    (Parcel.receiver_branch_id == user_branch_id) | 
                    (Parcel.created_by == current_user.id) |
                    (Parcel.receiver_email == current_user.email)
                )
        else:
            if direction == 'sent':
                query = query.filter(Parcel.created_by == current_user.id)
            elif direction == 'received':
                query = query.filter(Parcel.receiver_email == current_user.email)
            else:
                query = query.filter((Parcel.created_by == current_user.id) | (Parcel.receiver_email == current_user.email))
    else:
        # Superadmin: check if specific branch requested
        if branch_id and branch_id != 'all' and branch_id != '0':
            try:
                b_id = int(branch_id)
                if direction == 'sent':
                    query = query.filter(Parcel.sender_branch_id == b_id)
                elif direction == 'received':
                    query = query.filter(Parcel.receiver_branch_id == b_id)
                else: # both
                    query = query.filter((Parcel.sender_branch_id == b_id) | (Parcel.receiver_branch_id == b_id))
            except ValueError:
                pass
        else:
            if direction == 'sent':
                query = query.filter(Parcel.sender_branch_id != None)
            elif direction == 'received':
                query = query.filter(Parcel.receiver_branch_id != None)

    # 2. Status / Report Type Filter
    if report_type == 'pending':
        query = query.filter(Parcel.status.in_(['pending', 'in_transit']))
    elif report_type == 'delivered':
        query = query.filter(Parcel.status == 'delivered')
    elif report_type == 'failed':
        query = query.filter(Parcel.status == 'failed')
        
    # 3. Transfer Mode Filter
    if transfer_mode and transfer_mode != 'all':
        query = query.filter(Parcel.transfer_mode == transfer_mode)
        
    # 4. Courier Partner Filter
    if courier_partner and courier_partner != 'all':
        query = query.filter(Parcel.courier_partner.ilike(f'%{courier_partner}%'))
        
    # 5. Date Range Filtering
    if date_from:
        try:
            d_from = datetime.strptime(date_from, '%Y-%m-%d')
            query = query.filter(Parcel.dispatched_at >= d_from)
        except Exception:
            pass
    if date_to:
        try:
            d_to = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1) - timedelta(microseconds=1)
            query = query.filter(Parcel.dispatched_at <= d_to)
        except Exception:
            pass
            
    parcels = query.order_by(Parcel.dispatched_at.desc()).all()
    
    # Generate Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parcel Export Report"
    
    # Styled headers
    headers = [
        'Tracking Number', 'Status', 'Direction Type', 'Sender Name', 'Sender Email', 'Sender Phone', 'Sender Branch',
        'Receiver Name', 'Receiver Email', 'Receiver Phone', 'Destination Branch', 'Transfer Mode',
        'Courier Partner', 'Courier ID / AWB', 'Items / Description', 'Weight (kg)', 'Billing Amount (₹)', 
        'Special Instructions', 'Dispatched At', 'Expected Due Date', 'Delivered At'
    ]
    ws.append(headers)
    
    # Header Styling
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    header_fill = PatternFill(start_color="0D1B2A", end_color="0D1B2A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for p in parcels:
        sender_b = p.sender_branch.name if p.sender_branch else 'N/A'
        receiver_b = p.receiver_branch.name if p.receiver_branch else 'N/A'
        
        dir_type = "Sent" if (current_user.branch_id and p.sender_branch_id == current_user.branch_id) else "Received"
        
        ws.append([
            p.tracking_number,
            p.status_label,
            dir_type,
            p.sender_name,
            p.sender_email or '',
            p.sender_phone or '',
            sender_b,
            p.receiver_name,
            p.receiver_email or '',
            p.receiver_phone or '',
            receiver_b,
            p.transfer_mode_label,
            p.courier_partner or '—',
            p.courier_id or '—',
            p.description or '—',
            p.weight or 0.0,
            p.billing_amount or 0.0,
            p.special_instructions or '—',
            p.dispatched_at.strftime('%Y-%m-%d %H:%M') if p.dispatched_at else '—',
            p.due_date.strftime('%Y-%m-%d') if p.due_date else '—',
            p.delivered_at.strftime('%Y-%m-%d %H:%M') if p.delivered_at else '—'
        ])
        
    # Auto adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"parcel_report_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )
